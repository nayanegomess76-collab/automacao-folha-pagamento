import re
import unicodedata
import pdfplumber
from openpyxl import load_workbook


# =========================================================
# CAMINHOS
# =========================================================

ARQUIVO_PDF = r""
ARQUIVO_EXCEL = r""


# =========================================================
# FUNCOES AUXILIARES
# =========================================================

def normalizar_texto(texto):
    if texto is None:
        return ""

    texto = str(texto)
    texto = unicodedata.normalize("NFKD", texto)
    texto = texto.encode("ASCII", "ignore").decode("ASCII")
    texto = texto.upper()
    texto = re.sub(r"\s+", " ", texto)

    return texto.strip()


def valor_brasileiro(valor):
    if valor is None:
        return 0.0

    valor = str(valor).strip()

    if not valor:
        return 0.0

    valor = valor.replace(".", "")
    valor = valor.replace(",", ".")

    try:
        return float(valor)
    except:
        return 0.0


# =========================================================
# EXTRACAO DO NOME
# =========================================================

def extrair_nome(texto):

    texto = unicodedata.normalize("NFKD", texto)
    texto = texto.encode("ASCII", "ignore").decode("ASCII")

    for linha in texto.splitlines():

        linha = linha.strip()

        match = re.match(
            r"^\d+\s+(.+?)\s+\d{6}\s+1\s+1$",
            linha
        )

        if match:
            return match.group(1).strip().upper()

    return None


# =========================================================
# EXTRACAO DO VALOR LIQUIDO
# =========================================================

def extrair_valor_liquido(texto):

    texto_normalizado = normalizar_texto(texto)

    padroes = [
        r"VALOR LIQUIDO\s+([\d\.,]+)",
        r"VALOR LIQUIDO\s*[:\-]?\s*([\d\.,]+)"
    ]

    for padrao in padroes:

        match = re.search(
            padrao,
            texto_normalizado,
            re.IGNORECASE
        )

        if match:
            return valor_brasileiro(match.group(1))

    return 0.0


# =========================================================
# EXTRACAO DO FGTS
# =========================================================

def extrair_fgts(texto):

    texto_normalizado = normalizar_texto(texto)

    padrao = (
        r"SALARIO BASE\s+"
        r"SAL\. CONTR\. INSS\s+"
        r"BASE CALC\. FGTS\s+"
        r"F\.G\.T\.S DO MES\s+"
        r"BASE CALC\. IRRF\s+"
        r"FAIXA IRRF\s+"
        r"([\d\.,]+)\s+"
        r"([\d\.,]+)\s+"
        r"([\d\.,]+)\s+"
        r"([\d\.,]+)\s+"
        r"([\d\.,]+)"
    )

    match = re.search(
        padrao,
        texto_normalizado,
        re.IGNORECASE
    )

    if match:
        return valor_brasileiro(match.group(4))

    return 0.0


# =========================================================
# EXTRACAO DO INSS
# =========================================================

def extrair_inss(texto):

    texto_normalizado = normalizar_texto(texto)

    padrao = r"998\s+I\.N\.S\.S\.\s+([\d\.,]+)\s+([\d\.,]+)"

    match = re.search(
        padrao,
        texto_normalizado,
        re.IGNORECASE
    )

    if match:
        return valor_brasileiro(match.group(2))

    padrao_empregador = (
        r"843\s+INSS EMPREGADOR\s+"
        r"([\d\.,]+)\s+"
        r"([\d\.,]+)"
    )

    match = re.search(
        padrao_empregador,
        texto_normalizado,
        re.IGNORECASE
    )

    if match:
        return valor_brasileiro(match.group(2))

    return 0.0


# =========================================================
# LOCALIZAR COLUNAS
# =========================================================

def encontrar_coluna(ws, nome_coluna):

    nome_procurado = normalizar_texto(nome_coluna)

    for row in ws.iter_rows():

        for cell in row:

            if cell.value is None:
                continue

            valor = normalizar_texto(cell.value)

            if valor == nome_procurado:
                return cell.column

    return None


# =========================================================
# ABRINDO EXCEL
# =========================================================

print("=" * 70)
print("ABRINDO EXCEL")
print("=" * 70)

wb = load_workbook(ARQUIVO_EXCEL)

ws_folha = wb["FOLHA DE PGTO"]
ws_tributos = wb["TRIBUTOS"]


# =========================================================
# LOCALIZAR COLUNAS
# =========================================================

col_folha_colaborador = encontrar_coluna(
    ws_folha,
    "COLABORADOR"
)

col_folha_saldo = encontrar_coluna(
    ws_folha,
    "saldo 1"
)

col_tributos_colaborador = encontrar_coluna(
    ws_tributos,
    "COLABORADOR"
)

col_tributos_fgts = encontrar_coluna(
    ws_tributos,
    "FGTS"
)

col_tributos_inss = encontrar_coluna(
    ws_tributos,
    "INSS"
)


print()
print("COLUNAS ENCONTRADAS:")
print("FOLHA - COLABORADOR:", col_folha_colaborador)
print("FOLHA - SALDO 1:", col_folha_saldo)
print("TRIBUTOS - COLABORADOR:", col_tributos_colaborador)
print("TRIBUTOS - FGTS:", col_tributos_fgts)
print("TRIBUTOS - INSS:", col_tributos_inss)


# =========================================================
# VALIDAR COLUNAS
# =========================================================

if not col_folha_colaborador:
    raise Exception("Coluna COLABORADOR nao encontrada na FOLHA DE PGTO.")

if not col_folha_saldo:
    raise Exception("Coluna saldo 1 nao encontrada na FOLHA DE PGTO.")

if not col_tributos_colaborador:
    raise Exception("Coluna COLABORADOR nao encontrada em TRIBUTOS.")

if not col_tributos_fgts:
    raise Exception("Coluna FGTS nao encontrada em TRIBUTOS.")

if not col_tributos_inss:
    raise Exception("Coluna INSS nao encontrada em TRIBUTOS.")


# =========================================================
# MONTAR DICIONARIOS
# =========================================================

funcionarios_folha = {}

for linha in range(1, ws_folha.max_row + 1):

    valor = ws_folha.cell(
        linha,
        col_folha_colaborador
    ).value

    if valor is None:
        continue

    nome = normalizar_texto(valor)

    if nome == "COLABORADOR":
        continue

    if nome:
        funcionarios_folha[nome] = linha


funcionarios_tributos = {}

for linha in range(1, ws_tributos.max_row + 1):

    valor = ws_tributos.cell(
        linha,
        col_tributos_colaborador
    ).value

    if valor is None:
        continue

    nome = normalizar_texto(valor)

    if nome == "COLABORADOR":
        continue

    if nome:
        funcionarios_tributos[nome] = linha


print()
print("Funcionarios na FOLHA:", len(funcionarios_folha))
print("Funcionarios em TRIBUTOS:", len(funcionarios_tributos))


# =========================================================
# PROCESSANDO PDF
# =========================================================

print()
print("=" * 70)
print("PROCESSANDO PDF")
print("=" * 70)

paginas_processadas = 0
nao_encontrados_folha = []
nao_encontrados_tributos = []
paginas_erro = 0


with pdfplumber.open(ARQUIVO_PDF) as pdf:

    print()
    print("Paginas encontradas:", len(pdf.pages))

    for numero_pagina, pagina in enumerate(
        pdf.pages,
        start=1
    ):

        try:

            texto = pagina.extract_text()

            if not texto:

                print(
                    f"Pagina {numero_pagina}: texto nao encontrado"
                )

                paginas_erro += 1
                continue


            # =================================================
            # NOME
            # =================================================

            nome = extrair_nome(texto)

            if not nome:

                print(
                    f"Pagina {numero_pagina}: nome nao encontrado"
                )

                paginas_erro += 1
                continue


            print()
            print(
                f"Pagina {numero_pagina}: {nome}"
            )


            nome_normalizado = normalizar_texto(nome)


            # =================================================
            # VALORES
            # =================================================

            liquido = extrair_valor_liquido(texto)
            fgts = extrair_fgts(texto)
            inss = extrair_inss(texto)


            print(
                f"  Liquido: R$ {liquido:.2f}"
            )

            print(
                f"  FGTS: R$ {fgts:.2f}"
            )

            print(
                f"  INSS: R$ {inss:.2f}"
            )


            # =================================================
            # FOLHA DE PGTO
            # =================================================

            if nome_normalizado in funcionarios_folha:

                linha_folha = funcionarios_folha[
                    nome_normalizado
                ]

                ws_folha.cell(
                    linha_folha,
                    col_folha_saldo
                ).value = liquido

                print(
                    "  -> Saldo 1 atualizado"
                )

            else:

                print(
                    "  -> Nao encontrado na FOLHA DE PGTO"
                )

                nao_encontrados_folha.append(nome)


            # =================================================
            # TRIBUTOS
            # =================================================

            if nome_normalizado in funcionarios_tributos:

                linha_tributos = funcionarios_tributos[
                    nome_normalizado
                ]

                ws_tributos.cell(
                    linha_tributos,
                    col_tributos_fgts
                ).value = fgts

                ws_tributos.cell(
                    linha_tributos,
                    col_tributos_inss
                ).value = inss

                print(
                    "  -> FGTS e INSS atualizados"
                )

            else:

                print(
                    "  -> Nao encontrado em TRIBUTOS"
                )

                nao_encontrados_tributos.append(nome)


            paginas_processadas += 1


        except Exception as erro:

            print()
            print(
                f"Pagina {numero_pagina}: ERRO"
            )

            print(
                "Detalhes:",
                erro
            )

            paginas_erro += 1


# =========================================================
# SALVAR
# =========================================================

print()
print("=" * 70)
print("SALVANDO EXCEL")
print("=" * 70)

wb.save(ARQUIVO_EXCEL)

wb.close()

print()
print("Arquivo salvo.")
print("Caminho:")
print(ARQUIVO_EXCEL)


# =========================================================
# VERIFICAR O ARQUIVO REALMENTE SALVO
# =========================================================

print()
print("=" * 70)
print("VERIFICANDO ARQUIVO SALVO")
print("=" * 70)

wb_verificacao = load_workbook(ARQUIVO_EXCEL)

ws_folha_verificacao = wb_verificacao["FOLHA DE PGTO"]
ws_tributos_verificacao = wb_verificacao["TRIBUTOS"]


# =========================================================
# VERIFICAR ALGUNS FUNCIONARIOS
# =========================================================

funcionarios_teste = list(funcionarios_folha.keys())[:4]


print()
print("VERIFICACAO DO SALDO 1:")

for nome in funcionarios_teste:

    nome_normalizado = normalizar_texto(nome)

    linha = None

    for numero_linha in range(
        1,
        ws_folha_verificacao.max_row + 1
    ):

        valor = ws_folha_verificacao.cell(
            numero_linha,
            col_folha_colaborador
        ).value

        if normalizar_texto(valor) == nome_normalizado:

            linha = numero_linha
            break


    if linha:

        valor_salvo = ws_folha_verificacao.cell(
            linha,
            col_folha_saldo
        ).value

        print(
            f"{nome}: saldo 1 = {valor_salvo}"
        )

    else:

        print(
            f"{nome}: NAO ENCONTRADO"
        )


print()
print("VERIFICACAO DO FGTS E INSS:")

for nome in funcionarios_teste:

    nome_normalizado = normalizar_texto(nome)

    linha = None

    for numero_linha in range(
        1,
        ws_tributos_verificacao.max_row + 1
    ):

        valor = ws_tributos_verificacao.cell(
            numero_linha,
            col_tributos_colaborador
        ).value

        if normalizar_texto(valor) == nome_normalizado:

            linha = numero_linha
            break


    if linha:

        valor_fgts = ws_tributos_verificacao.cell(
            linha,
            col_tributos_fgts
        ).value

        valor_inss = ws_tributos_verificacao.cell(
            linha,
            col_tributos_inss
        ).value

        print(
            f"{nome}: FGTS = {valor_fgts} | INSS = {valor_inss}"
        )

    else:

        print(
            f"{nome}: NAO ENCONTRADO"
        )


wb_verificacao.close()


# =========================================================
# RESULTADO FINAL
# =========================================================

print()
print("=" * 70)
print("PROCESSAMENTO FINALIZADO")
print("=" * 70)

print(
    "Paginas processadas:",
    paginas_processadas
)

print(
    "Paginas com erro:",
    paginas_erro
)

print(
    "Nao encontrados na FOLHA:",
    len(nao_encontrados_folha)
)

print(
    "Nao encontrados em TRIBUTOS:",
    len(nao_encontrados_tributos)
)


if nao_encontrados_folha:

    print()
    print("Nomes nao encontrados na FOLHA:")

    for nome in nao_encontrados_folha:
        print("-", nome)


if nao_encontrados_tributos:

    print()
    print("Nomes nao encontrados em TRIBUTOS:")

    for nome in nao_encontrados_tributos:
        print("-", nome)


print()
print("Excel atualizado com sucesso.")
print("O IRRF nao foi alterado.")